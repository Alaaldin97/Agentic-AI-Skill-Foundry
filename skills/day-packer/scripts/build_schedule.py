#!/usr/bin/env python3
"""Build an Excel timetable and an .ics calendar file from a schedule JSON.

Usage: python3 build_schedule.py schedule.json [--out-dir DIR] [--name NAME]
"""
import argparse
import json
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FMT = "%Y-%m-%dT%H:%M"

CATEGORY_COLORS = {
    "deep": "C9DAF8",
    "routine": "E2EFDA",
    "physical": "FCE4D6",
    "errand": "FFF2CC",
    "social": "EAD1DC",
    "travel": "D9D9D9",
    "sleep": "BDD7EE",
    "break": "EDEDED",
    "fixed": "F4CCCC",
}


def parse_blocks(data):
    blocks = []
    for i, b in enumerate(data.get("blocks", [])):
        try:
            start = datetime.strptime(b["start"], FMT)
            end = datetime.strptime(b["end"], FMT)
        except (KeyError, ValueError) as e:
            sys.exit(f"Block {i} ({b.get('title', '?')}): bad or missing start/end ({e}). Use YYYY-MM-DDTHH:MM.")
        if end <= start:
            sys.exit(f"Block {i} ({b.get('title', '?')}): end {b['end']} is not after start {b['start']}.")
        blocks.append({**b, "_start": start, "_end": end})
    blocks.sort(key=lambda x: x["_start"])
    return blocks


def check_overlaps(blocks):
    warnings = []
    for a, b in zip(blocks, blocks[1:]):
        if b["_start"] < a["_end"]:
            stacked = "(stacked)" in a.get("title", "") or "(stacked)" in b.get("title", "")
            if not stacked:
                warnings.append(f"OVERLAP: '{a['title']}' ends {a['end']} but '{b['title']}' starts {b['start']}")
    return warnings


def write_xlsx(blocks, did_not_fit, title, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Date", "Start", "End", "Minutes", "Task", "Where", "Category", "Notes"]
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=3, column=c).font = Font(bold=True)
    total = 0
    for b in blocks:
        mins = int((b["_end"] - b["_start"]).total_seconds() // 60)
        total += mins
        end_label = b["_end"].strftime("%H:%M")
        if b["_end"].date() != b["_start"].date():
            end_label += " (+1d)" if (b["_end"].date() - b["_start"].date()).days == 1 else f" ({b['_end'].strftime('%a %d')})"
        ws.append([
            b["_start"].strftime("%a %d %b"),
            b["_start"].strftime("%H:%M"),
            end_label,
            mins,
            b.get("title", ""),
            b.get("location", ""),
            b.get("category", ""),
            b.get("notes", ""),
        ])
        color = CATEGORY_COLORS.get(str(b.get("category", "")).lower())
        if color:
            fill = PatternFill("solid", start_color=color, end_color=color)
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).fill = fill
    ws.append([])
    ws.append(["", "", "Total", total, f"{total // 60}h {total % 60}m scheduled"])
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    widths = [14, 8, 14, 9, 40, 18, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A4"

    ws2 = wb.create_sheet("Didn't fit")
    ws2.append(["Task", "Reason"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    for item in did_not_fit:
        ws2.append([item.get("title", ""), item.get("reason", "")])
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 50
    wb.save(path)


def ics_escape(text):
    return (str(text).replace("\\", "\\\\").replace(";", "\\;")
            .replace(",", "\\,").replace("\n", "\\n"))


def fold(line):
    """Fold lines longer than 75 octets per RFC 5545."""
    out, cur = [], ""
    for ch in line:
        if len((cur + ch).encode("utf-8")) > 75:
            out.append(cur)
            cur = " " + ch
        else:
            cur += ch
    out.append(cur)
    return "\r\n".join(out)


def write_ics(blocks, title, path):
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//day-packer//EN",
        "CALSCALE:GREGORIAN",
        f"X-WR-CALNAME:{ics_escape(title)}",
    ]
    for b in blocks:
        desc_parts = [p for p in [b.get("category", ""), b.get("notes", "")] if p]
        lines += [
            "BEGIN:VEVENT",
            f"UID:{uuid.uuid4()}@day-packer",
            f"DTSTAMP:{stamp}",
            # Floating local time (no TZID): shows at the same clock time in the user's calendar timezone
            f"DTSTART:{b['_start'].strftime('%Y%m%dT%H%M%S')}",
            f"DTEND:{b['_end'].strftime('%Y%m%dT%H%M%S')}",
            f"SUMMARY:{ics_escape(b.get('title', ''))}",
        ]
        if b.get("location"):
            lines.append(f"LOCATION:{ics_escape(b['location'])}")
        if desc_parts:
            lines.append(f"DESCRIPTION:{ics_escape(' | '.join(desc_parts))}")
        lines.append("END:VEVENT")
    lines.append("END:VCALENDAR")
    Path(path).write_text("\r\n".join(fold(l) for l in lines) + "\r\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("schedule_json")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--name", default=None, help="Base filename (default: from title)")
    args = ap.parse_args()

    data = json.loads(Path(args.schedule_json).read_text(encoding="utf-8"))
    title = data.get("title", "Schedule")
    blocks = parse_blocks(data)
    if not blocks:
        sys.exit("No blocks in schedule.")

    warnings = check_overlaps(blocks)
    for w in warnings:
        print(w)

    name = args.name or re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_").lower() or "schedule"
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    xlsx_path = out / f"{name}.xlsx"
    ics_path = out / f"{name}.ics"
    write_xlsx(blocks, data.get("did_not_fit", []), title, xlsx_path)
    write_ics(blocks, title, ics_path)
    print(f"Wrote {xlsx_path}")
    print(f"Wrote {ics_path}")
    if warnings:
        print(f"{len(warnings)} overlap warning(s) - fix before delivering.")


if __name__ == "__main__":
    main()
